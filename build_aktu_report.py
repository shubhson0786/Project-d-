from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from pathlib import Path

ROOT = Path(r"C:\Users\shubham sonkar\Desktop\project d")
OUT = ROOT / "Stashflow_AKTU_Project_Report_Completed.docx"
COLLEGE_NAME = "DR AMBEDKAR INSTITUTE OF TECHNOLOGY FOR DIVYANGJAN KANPUR"
GUIDE_NAME = "Sarvesh Sir"
MEMBERS = [
    ("Shubham Sonkar", "2401661530066"),
    ("Mradul Singh", "2401661530035"),
    ("Divya Tiwari", "2401661530018"),
    ("Ashi Shukla", "2401661530013"),
]

doc = Document()

def font(run, size=12, bold=False, italic=False, name="Times New Roman", color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    return run

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(12)
normal.paragraph_format.line_spacing = 1.5
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

for style_name, size, before, after in [("Title", 18, 0, 18), ("Heading 1", 16, 18, 12), ("Heading 2", 14, 14, 8), ("Heading 3", 12, 10, 6)]:
    st = styles[style_name]
    st.font.name = "Times New Roman"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

if "Caption" in styles:
    cap = styles["Caption"]
else:
    cap = styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
cap.font.name = "Times New Roman"
cap._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
cap._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
cap.font.size = Pt(10)
cap.font.italic = True
cap.font.color.rgb = RGBColor(0, 0, 0)
cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_after = Pt(10)

for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(1.25)
    section.bottom_margin = Inches(1.05)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(0.85)

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")

def set_table_borders(table, color="D9D9D9", size="6"):
    tblPr = table._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders"); tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}"); borders.append(el)
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), size); el.set(qn("w:color"), color)

def add_table(headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for j, h in enumerate(headers):
        c = table.rows[0].cells[j]
        set_cell_shading(c, "1F4E78"); set_cell_margins(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(h), 10, True, color="FFFFFF")
        if widths: c.width = Inches(widths[j])
    trPr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.append(tbl_header)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, value in enumerate(row):
            c = cells[j]; set_cell_margins(c); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2: set_cell_shading(c, "F2F6FA")
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 25 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
            font(p.add_run(str(value)), 9.5)
            if widths: c.width = Inches(widths[j])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_caption(text):
    doc.add_paragraph(text, style="Caption")

def add_p(text="", bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        font(p.add_run(bold_lead), 12, True)
        font(p.add_run(text[len(bold_lead):]), 12)
    else:
        font(p.add_run(text), 12)
    return p

def add_bullets(items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        p.paragraph_format.line_spacing = 1.35
        font(p.add_run(item), 12)

def add_numbered(items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.line_spacing = 1.35
        font(p.add_run(item), 12)

def centered(text, size=12, bold=False, space_after=6):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(space_after)
    font(p.add_run(text), size, bold)
    return p

def field(run, instruction):
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = instruction
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, sep, end])

def set_page_number_format(section, fmt="decimal", start=1):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None: pg = OxmlElement("w:pgNumType"); sectPr.append(pg)
    pg.set(qn("w:fmt"), fmt); pg.set(qn("w:start"), str(start))
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    for child in list(p._p):
        p._p.remove(child)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(); font(r, 10); field(r, "PAGE")

def new_page():
    doc.add_page_break()

def chapter(number, title):
    p = doc.add_paragraph(style="Heading 1"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    prefix = f"APPENDIX {number}" if isinstance(number, str) and number.isalpha() else f"CHAPTER {number}"
    font(p.add_run(f"{prefix}\n{title.upper()}"), 16, True)

# Cover page
sec = doc.sections[0]
sec.different_first_page_header_footer = True
centered("A PROJECT REPORT", 16, True, 6)
centered("ON", 12, True, 4)
centered("STASHFLOW", 22, True, 2)
centered("Web Based Investment and Trading Dashboard", 14, True, 12)
centered("Submitted in partial fulfilment of the requirements for the award of the degree of", 12, False, 8)
centered("BACHELOR OF TECHNOLOGY", 15, True, 2)
centered("IN", 11, True, 2)
centered("COMPUTER SCIENCE AND ENGINEERING\n(ARTIFICIAL INTELLIGENCE AND MACHINE LEARNING)", 13, True, 10)
centered("Submitted by", 12, True, 8)
add_table(["Student Name", "University Roll Number"], [[name, roll] for name, roll in MEMBERS], [3.2, 3.2])
centered("Under the guidance of", 12, True, 6)
centered(GUIDE_NAME, 12, True, 2)
centered("Project Guide", 11, False, 8)
centered("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING\n(ARTIFICIAL INTELLIGENCE AND MACHINE LEARNING)", 10, True, 4)
centered("_______________________________________________", 12, True, 4)
centered(COLLEGE_NAME, 10, True, 4)
centered("DR A P J ABDUL KALAM TECHNICAL UNIVERSITY UTTAR PRADESH LUCKNOW", 11, True, 4)
centered("SESSION 2026 2027", 12, True, 0)

# Preliminary section with Roman numbering
pre = doc.add_section(WD_SECTION.NEW_PAGE)
pre.page_width = Inches(8.27); pre.page_height = Inches(11.69)
pre.top_margin = Inches(1.25); pre.bottom_margin = Inches(1.05); pre.left_margin = Inches(1.5); pre.right_margin = Inches(0.85)
set_page_number_format(pre, "lowerRoman", 1)

centered("BONAFIDE CERTIFICATE", 16, True, 24)
add_p("This is to certify that the project report entitled “Stashflow: Web Based Investment and Trading Dashboard” is a bonafide record of the project work carried out by the student or project team named below under my supervision and guidance. The work is submitted to Dr A P J Abdul Kalam Technical University, Uttar Pradesh, Lucknow, in partial fulfilment of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering (Artificial Intelligence and Machine Learning) during the academic session 2026–2027.")
add_p("To the best of my knowledge, the work presented in this report has not been submitted elsewhere for the award of any other degree or diploma. The implementation, observations, limitations, and results stated in the report have been examined for academic submission.")
doc.add_paragraph("\n")
add_table(["Student Name", "University Roll Number", "Signature"], [[name, roll, "________________"] for name, roll in MEMBERS], [2.4, 2.4, 1.6])
doc.add_paragraph("\n\n")
add_table(["Project Guide", "Head of Department"], [[f"Name: {GUIDE_NAME}\nSignature: _____________________\nDate: _________________________", "Name: ________________________\nSignature: _____________________\nDate: _________________________"]], [3.2, 3.2])
centered("Department Seal", 10, False, 0)

new_page(); centered("STUDENT DECLARATION", 16, True, 24)
add_p("We hereby declare that the project report entitled “Stashflow: Web Based Investment and Trading Dashboard” is an original account of the work undertaken by us under the supervision of the project guide named in this report. The source code, design decisions, testing evidence, and technical discussion have been prepared for this academic project. Wherever the ideas, documentation, or software of other authors and organisations have been used, suitable acknowledgement has been provided in the references.")
add_p("We further declare that this report has not been submitted, either in full or in part, for the award of any degree or diploma at this or any other university. We accept responsibility for the correctness of the reported work and understand that the institution may verify the report for originality and technical accuracy.")
doc.add_paragraph("\n\n")
add_table(["Student Name", "University Roll Number", "Signature and Date"], [[name, roll, "________________________"] for name, roll in MEMBERS], [2.3, 2.3, 1.8])

new_page(); centered("ACKNOWLEDGEMENT", 16, True, 24)
add_p("We express our sincere gratitude to our project guide for continuous guidance, technical suggestions, and careful review during the development of Stashflow. The guidance helped us convert an initial idea into a structured web application and encouraged us to evaluate the system not only for visible functionality but also for data validation, security, reliability, and deployment behaviour.")
add_p("We thank the Head of the Department of Computer Science and Engineering (Artificial Intelligence and Machine Learning) and all faculty members for providing the academic environment and resources required for this work. We also acknowledge Dr A P J Abdul Kalam Technical University, Uttar Pradesh, for the curriculum that motivated the application of software engineering, database management, web technology, and information security concepts in an integrated project.")
add_p("Finally, we thank our family members, classmates, and test users for their encouragement and practical feedback. Their observations helped improve navigation, form validation, responsive presentation, and the clarity of the investment dashboard.")
doc.add_paragraph("\n\n")
add_table(["Student Name", "Signature"], [[name, "____________________________"] for name, _ in MEMBERS], [3.2, 3.2])

new_page(); centered("ABSTRACT", 16, True, 18)
add_p("Stashflow is a full-stack academic prototype that demonstrates how a modern web application can present investment information, accept simulated trading orders, manage a virtual fund balance, and integrate a payment gateway in test mode. The system separates the public landing experience from the trading dashboard. The landing application presents product information, pricing, support content, and a combined registration and login interface. The dashboard provides watchlist search, portfolio summaries, holdings, positions, order entry, order history, fund deposits, withdrawals, and transaction history.")
add_p("The user interfaces are implemented in React and use client-side routing for page navigation. MongoDB stores user registrations, holdings, positions, orders, fund accounts, transaction records, and Razorpay order references. Node.js serverless API functions perform database operations, validate requests, hash account passwords with bcrypt, create Razorpay orders, and verify payment signatures using HMAC SHA-256. Vercel hosts the landing site and dashboard as separate deployments, while Razorpay Test Mode provides a safe simulated payment experience without moving real money.")
add_p("The implemented prototype successfully builds and deploys, demonstrates responsive investment workflows, and prevents duplicate crediting of the same verified payment through persistent payment-order state and database transactions. Review and testing also identify important limitations: dashboard APIs require complete session enforcement, financial records require per-user ownership, automated tests are limited, dependency upgrades are needed, and payment webhooks should be added before any live use. The work therefore meets its academic objective as a demonstrable prototype while providing a clear roadmap for production hardening.")

new_page(); centered("TABLE OF CONTENTS", 16, True, 12)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run(); font(r, 12); field(r, 'TOC \\o "1-3" \\h \\z \\u')
add_p("Note: In Microsoft Word, select the table and choose Update Field if page numbers do not refresh automatically.")

new_page(); centered("LIST OF FIGURES", 16, True, 12)
add_table(["Figure", "Title"], [
    ["4.1", "Logical architecture of Stashflow"], ["4.2", "Authentication request flow"],
    ["4.3", "Razorpay test payment sequence"], ["5.1", "Dashboard navigation structure"],
    ["5.2", "Fund transaction state transition"],
], [1.1, 5.3])
centered("LIST OF TABLES", 16, True, 12)
add_table(["Table", "Title"], [
    ["2.1", "Comparison of the existing and proposed approaches"], ["3.1", "Functional requirements"],
    ["3.2", "Non-functional requirements"], ["3.3", "Hardware and software requirements"],
    ["4.1", "Primary data entities"], ["5.1", "Main application modules"],
    ["6.1", "Functional test cases"], ["6.2", "Build and verification results"],
    ["7.1", "Risk and mitigation summary"],
], [1.1, 5.3])

new_page(); centered("LIST OF ABBREVIATIONS", 16, True, 12)
add_table(["Abbreviation", "Meaning"], [
    ["API", "Application Programming Interface"], ["CRUD", "Create Read Update Delete"],
    ["HMAC", "Hash based Message Authentication Code"], ["HTTP", "Hypertext Transfer Protocol"],
    ["INR", "Indian Rupee"], ["JSON", "JavaScript Object Notation"],
    ["JWT", "JSON Web Token"], ["NoSQL", "Non relational database model"],
    ["SPA", "Single Page Application"], ["UI", "User Interface"],
    ["UPI", "Unified Payments Interface"], ["UX", "User Experience"],
], [1.5, 4.9])

# Main section
main = doc.add_section(WD_SECTION.NEW_PAGE)
main.page_width = Inches(8.27); main.page_height = Inches(11.69)
main.top_margin = Inches(1.25); main.bottom_margin = Inches(1.05); main.left_margin = Inches(1.5); main.right_margin = Inches(0.85)
set_page_number_format(main, "decimal", 1)

doc.add_paragraph("CHAPTER 1\nINTRODUCTION", style="Heading 1").alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_heading("1.1 Background", level=2)
add_p("Digital investing has increased the need for interfaces that explain portfolio information without forcing users to move between disconnected tools. A typical investor checks a watchlist, evaluates holdings and open positions, places an order, reviews available cash, and tracks account transactions. When these activities are represented through unrelated pages or static records, users lose context and may make avoidable errors. An integrated dashboard can reduce this friction by presenting common actions through one consistent navigation model.")
add_p("Stashflow was developed as an academic full-stack web project to explore this integration. The application combines a public information website, account registration and login, a trading-style dashboard, persistent MongoDB records, and a Razorpay Test Mode payment flow. The project does not execute trades on a stock exchange and does not move real money. It models the interaction and data-processing patterns required by such a system so that front-end, API, database, security, and deployment concepts can be evaluated together.")

doc.add_heading("1.2 Problem Statement", level=2)
add_p("Many student-built investment interfaces remain limited to static pages. They may display sample prices but do not preserve orders, connect user input to a database, or demonstrate how a payment provider should be invoked and verified. Other prototypes focus on backend operations but provide little attention to usability, responsive layouts, input validation, or failure messages. The project problem is therefore to design and implement a coherent web application that demonstrates the principal workflows of an investment dashboard while maintaining clear separation between client-side presentation and server-side trust decisions.")

doc.add_heading("1.3 Objectives", level=2)
add_bullets([
    "Develop a responsive public website that communicates the purpose, products, pricing, support options, and account entry points of Stashflow.",
    "Provide registration and login forms with server-side password hashing and persistent user records.",
    "Present a dashboard containing summary, watchlist, holdings, positions, orders, funds, and applications sections.",
    "Validate and store simulated buy and sell orders through API functions and MongoDB.",
    "Integrate Razorpay Standard Checkout in Test Mode, create orders on the server, and verify returned payment signatures before crediting the virtual balance.",
    "Deploy the landing application and dashboard on Vercel and verify that their production builds and APIs are reachable.",
    "Evaluate the implementation honestly and document the controls still required before live financial use.",
])

doc.add_heading("1.4 Scope", level=2)
add_p("The project scope covers web-based interaction for an educational investment dashboard. It includes a React landing application, a separate React dashboard, Node.js or Vercel serverless API functions, MongoDB persistence, chart presentation, authentication registration and login, order recording, and simulated fund management. Market quotations and watchlist entries are demonstration data. The system does not connect to a stock exchange, depository participant, bank payout API, KYC service, or regulatory reporting service. Razorpay is configured only with test credentials.")

doc.add_heading("1.5 Development Method", level=2)
add_p("An iterative development method was followed. Initial pages established routing and visual hierarchy. Data-oriented dashboard modules were then connected to MongoDB APIs. Form validation and error states were added around order entry and funds management. The payment workflow was implemented as a server-created Razorpay order followed by client checkout and server-side signature verification. Finally, both React applications were built and deployed, and the completed system was reviewed for security, data ownership, dependency, and test-coverage limitations.")

doc.add_heading("1.6 Organisation of the Report", level=2)
add_p("Chapter 2 reviews related concepts and compares the proposed approach with common static prototypes. Chapter 3 defines requirements and feasibility. Chapter 4 presents the architecture, data design, and important request sequences. Chapter 5 explains implementation details. Chapter 6 describes testing and observed results. Chapter 7 analyses security, limitations, and risk controls. Chapter 8 concludes the work and identifies future enhancements. The appendices provide setup instructions, API summaries, and sample viva questions.")

chapter(2, "Literature Review and Existing System")
doc.add_heading("2.1 Web Based Investment Platforms", level=2)
add_p("Web investment platforms are information-intensive systems. Their interfaces must place price movement, portfolio value, order state, and available funds in a hierarchy that supports quick interpretation. Single-page applications are useful in this context because client-side routing and reusable components can update a selected area without reloading the full site. React applies a component model in which state changes drive interface updates, while chart libraries convert numerical series into visual summaries.")
add_p("A production platform normally separates public content, identity management, trading operations, payments, and reporting into controlled services. Stashflow uses a simplified form of this separation: a public landing application, an account API, a dashboard application, and dashboard-specific serverless functions. This design makes the project easier to demonstrate, although the review in Chapter 7 explains why cross-domain session enforcement and per-user data ownership must be strengthened.")

doc.add_heading("2.2 Authentication and Password Storage", level=2)
add_p("Authentication establishes the identity associated with a request. Passwords should never be stored as readable text. Instead, a slow password-hashing function should generate a salted hash that can be compared during login. Stashflow uses bcrypt for this purpose. A successful login also creates an HTTP-only cookie, which prevents normal client-side JavaScript from reading the session value. However, secure authentication is complete only when each protected API validates the session and applies the resulting user identifier to every data query.")

doc.add_heading("2.3 Database Persistence", level=2)
add_p("MongoDB represents application data as documents grouped into collections. This model is suitable for a prototype containing users, holdings, positions, orders, fund accounts, transaction arrays, and payment orders. Mongoose supplies schema validation, model operations, and transaction support. The project uses atomic update operators for balance changes and a database transaction when marking a Razorpay order as paid and crediting funds. These choices reduce inconsistent intermediate states.")

doc.add_heading("2.4 Payment Gateway Integration", level=2)
add_p("A payment gateway integration must keep the key secret on the server. The browser may receive a publishable key identifier and provider order identifier, but order creation and signature verification remain server responsibilities. Razorpay Standard Checkout returns a payment identifier, order identifier, and signature after a successful test interaction. Stashflow recomputes an HMAC SHA-256 digest from the stored order relationship and compares it with the returned signature. Test Mode simulates the flow and does not deduct or settle real money.")

doc.add_heading("2.5 Existing and Proposed Approaches", level=2)
add_table(["Criterion", "Common static prototype", "Stashflow prototype"], [
    ["Navigation", "Separate static pages", "Client-side routed landing site and dashboard"],
    ["Data persistence", "Hard-coded arrays only", "MongoDB-backed users, orders, funds, holdings and positions"],
    ["Authentication", "Browser-only flag", "Bcrypt password hashes and HTTP-only session cookie"],
    ["Order workflow", "Button without persistence", "Validated API request stored as an order document"],
    ["Payments", "Balance modified directly", "Razorpay test order plus server signature verification"],
    ["Deployment", "Localhost demonstration", "Separate Vercel production deployments"],
    ["Known limitation", "No security model", "Security model identified but not yet fully enforced on dashboard APIs"],
], [1.5, 2.35, 2.55])
add_caption("Table 2.1 Comparison of the existing and proposed approaches")

chapter(3, "Requirement Analysis and Feasibility")
doc.add_heading("3.1 Stakeholders", level=2)
add_p("The principal stakeholder is a student or evaluator who needs to explore the complete project workflow. A registered user interacts with the landing site and dashboard. The project team develops and maintains the source code and deployment configuration. The academic guide evaluates scope, design, implementation, results, and originality. External technical services include MongoDB, Vercel, and Razorpay Test Mode.")

doc.add_heading("3.2 Functional Requirements", level=2)
add_table(["ID", "Requirement", "Acceptance condition"], [
    ["FR-01", "Navigate public pages", "Home, About, Products, Pricing and Support routes render"],
    ["FR-02", "Register account", "Valid new user is stored with a password hash"],
    ["FR-03", "Log in", "Registered email and password return an authenticated response"],
    ["FR-04", "Search watchlist", "Visible symbols are filtered by entered text"],
    ["FR-05", "View portfolio", "Holdings and positions are obtained from APIs and summarised"],
    ["FR-06", "Place order", "Valid buy or sell order is stored and returned"],
    ["FR-07", "View order history", "Stored orders are displayed newest first"],
    ["FR-08", "Add funds", "Verified Razorpay test payment credits the virtual balance once"],
    ["FR-09", "Withdraw funds", "Valid amount reduces balance only when sufficient cash exists"],
    ["FR-10", "Review transactions", "Recent deposits and withdrawals display with date and method"],
], [0.7, 2.0, 3.7])
add_caption("Table 3.1 Functional requirements")

doc.add_heading("3.3 Non Functional Requirements", level=2)
add_table(["Category", "Requirement"], [
    ["Usability", "Consistent navigation, readable monetary formatting, responsive dialogs and actionable error messages"],
    ["Performance", "Reuse database connections in serverless functions and avoid full-page reloads during dashboard navigation"],
    ["Security", "Hash passwords, keep secrets server-side, verify payment signatures and validate all mutable input"],
    ["Reliability", "Use atomic updates and transaction-protected payment crediting to avoid partial writes"],
    ["Maintainability", "Separate UI components, shared API configuration, data models and environment settings"],
    ["Portability", "Run in a standards-based browser and deploy through a cloud platform using environment variables"],
], [1.5, 4.9])
add_caption("Table 3.2 Non-functional requirements")

doc.add_heading("3.4 Hardware and Software Requirements", level=2)
add_table(["Type", "Minimum or selected requirement"], [
    ["Development computer", "Dual-core processor, 8 GB RAM, 2 GB free storage"],
    ["Client device", "Desktop or mobile device with a modern JavaScript-enabled browser"],
    ["Operating system", "Windows 10 or later for development; cloud runtime for deployment"],
    ["Frontend", "React, React Router, CSS, Material UI, Chart.js and Axios"],
    ["Backend", "Node.js, Express for the standalone service, and Vercel serverless functions"],
    ["Database", "MongoDB with Mongoose object modelling"],
    ["Payments", "Razorpay Node SDK and Standard Checkout in Test Mode"],
    ["Hosting", "Vercel deployments for landing site and dashboard"],
], [1.7, 4.7])
add_caption("Table 3.3 Hardware and software requirements")

doc.add_heading("3.5 Feasibility Analysis", level=2)
add_p("Technical feasibility is supported by mature JavaScript libraries, a managed document database, and serverless deployment. The team can develop and test the principal workflow using commonly available computers and browsers. Operational feasibility follows from familiar navigation and direct form feedback. Economic feasibility is favourable for academic use because the selected tools and cloud services offer free development or test tiers. Schedule feasibility is improved by dividing the work into interface, API, database, payment, and deployment milestones.")

chapter(4, "System Design")
doc.add_heading("4.1 Architectural Overview", level=2)
add_p("Stashflow uses a layered web architecture. The presentation layer contains two React applications. The service layer contains authentication and dashboard API functions. The persistence layer stores documents in MongoDB. Razorpay operates as an external test payment provider, and Vercel supplies build, routing, and serverless execution for the deployed applications.")
add_table(["Public client", "Application services", "Data and external services"], [
    ["Landing React application\nRegistration and login", "Authentication API\nInput validation\nPassword hashing", "MongoDB users collection"],
    ["Dashboard React application\nPortfolio and funds UI", "Orders, holdings, positions and funds APIs\nRazorpay order and verification APIs", "MongoDB trading collections\nRazorpay Test Mode"],
], [2.0, 2.2, 2.2])
add_caption("Figure 4.1 Logical architecture of Stashflow")

doc.add_heading("4.2 Component Design", level=2)
add_p("The landing application contains reusable page sections for navigation, home content, product descriptions, pricing, support, account registration, and login. Its authentication form switches between login and signup modes by reading the query string. The dashboard contains a top bar, menu, watchlist, routed content area, reusable chart components, and an order-entry dialog supplied through React context. Each data-oriented page requests JSON from a relative /api endpoint.")

doc.add_heading("4.3 Data Design", level=2)
add_table(["Entity", "Important fields", "Purpose"], [
    ["User", "name, email, passwordHash, timestamps", "Stores registered account credentials"],
    ["Order", "name, qty, price, mode, timestamps", "Records simulated buy and sell orders"],
    ["Holding", "name, qty, avg, price, net, day", "Represents portfolio holdings"],
    ["Position", "product, name, qty, avg, price, net, day", "Represents open trading positions"],
    ["FundAccount", "accountKey, balance, usedMargin, transactions", "Maintains virtual account cash and history"],
    ["PaymentOrder", "razorpayOrderId, amount, currency, status, paymentId", "Tracks gateway orders and duplicate credit protection"],
], [1.35, 2.7, 2.35])
add_caption("Table 4.1 Primary data entities")

doc.add_heading("4.4 Authentication Flow", level=2)
add_numbered([
    "The user selects Login or Sign up on the shared account page.",
    "The browser validates basic password and confirmation requirements and posts JSON to the authentication API.",
    "For registration, the API checks whether the email already exists and hashes the password with bcrypt before storing the user.",
    "For login, the API fetches the user and compares the supplied password with the stored bcrypt hash.",
    "The API returns public profile fields and sets a signed HTTP-only cookie with a seven-day lifetime.",
    "The browser stores the display profile and redirects to the dashboard deployment.",
])
add_caption("Figure 4.2 Authentication request flow")

doc.add_heading("4.5 Payment Sequence", level=2)
add_numbered([
    "The user enters a deposit amount and submits the Add Funds form.",
    "The dashboard API validates the amount and creates an INR order using the Razorpay server SDK.",
    "The provider order identifier and public key identifier are returned to Standard Checkout.",
    "After a simulated successful payment, Checkout returns the payment identifier, order identifier, and signature.",
    "The verification API recomputes the HMAC SHA-256 signature with the server secret.",
    "Within a MongoDB transaction, the payment order is marked paid and the stored order amount is credited exactly once.",
    "The updated fund account and transaction history are returned to the dashboard.",
])
add_caption("Figure 4.3 Razorpay test payment sequence")

doc.add_heading("4.6 Validation and Error Design", level=2)
add_p("Validation is applied at both presentation and API layers. Order quantity must be a positive whole number, price must be positive, order mode must be BUY or SELL, and deposit or withdrawal amounts must be finite and at least one rupee. The API restricts deposits to Razorpay Checkout, verifies sufficient balance for withdrawals, rejects unsupported HTTP methods, and returns structured JSON errors. The interface presents these messages near the relevant form and disables submission while a request is in progress.")

chapter(5, "Implementation")
doc.add_heading("5.1 Project Structure", level=2)
add_p("The workspace contains three Node.js packages. The frontend directory contains the public React site and authentication API. The dashboard directory contains the trading React application and its serverless APIs. The backend directory contains an Express service with equivalent holdings, positions, and order endpoints for local or alternative hosting. Environment-specific values such as MongoDB connection strings, Razorpay credentials, and the session secret are supplied through environment variables rather than browser source code.")

doc.add_heading("5.2 Main Modules", level=2)
add_table(["Module", "Implemented responsibility"], [
    ["Landing site", "Product communication, responsive navigation, pricing, support, registration and login"],
    ["Authentication API", "Email normalization, duplicate detection, bcrypt hashing, password comparison and cookie creation"],
    ["Watchlist", "Static demonstration symbols, price movement, search and buy or sell actions"],
    ["Portfolio", "Holdings and positions retrieval, profit calculation and chart preparation"],
    ["Orders", "Validated simulated order submission and reverse chronological order history"],
    ["Funds", "Virtual balance, withdrawals, recent transactions and projected balance"],
    ["Razorpay integration", "Server order creation, Standard Checkout, signature verification and idempotent crediting"],
    ["Deployment", "Vercel builds, SPA rewrites, serverless APIs and environment configuration"],
], [1.7, 4.7])
add_caption("Table 5.1 Main application modules")

doc.add_heading("5.3 Frontend Implementation", level=2)
add_p("React functional components and hooks manage interface state. useState stores form values, loading indicators, error messages, modal state, and API results. useEffect loads data when portfolio pages mount and attaches or removes Escape-key handlers for dialogs. useMemo calculates portfolio totals, order values, and projected balances without repeating work unnecessarily. React Router maps browser paths to pages while retaining the single-page application shell.")
add_p("The interface applies Indian currency formatting with the en-IN locale. Reusable CSS classes create responsive cards, transaction rows, modal backdrops, form controls, and mobile breakpoints. Material UI icons support watchlist actions, while Chart.js renders doughnut and vertical bar charts for portfolio information.")

doc.add_heading("5.4 Authentication Implementation", level=2)
add_p("The account page presents Login and Sign up as adjacent tabs. Registration requests include the name, normalized email, and password. The server validates the email format and minimum password length, checks uniqueness, and stores only a bcrypt hash. Login retrieves the normalized email and uses bcrypt.compare to validate the secret. The API returns only the name and email to the browser. A signed session value is delivered in a Secure, HTTP-only, SameSite=Lax cookie.")

doc.add_heading("5.5 Dashboard Navigation", level=2)
add_table(["Navigation route", "Rendered content"], [
    ["/", "Account summary and overview"], ["/orders", "Saved order history"],
    ["/holdings", "Portfolio holdings and allocation chart"], ["/positions", "Open positions and profit values"],
    ["/funds", "Balance cards, deposit or withdrawal dialog and transaction history"], ["/apps", "Associated application placeholder"],
], [2.0, 4.4])
add_caption("Figure 5.1 Dashboard navigation structure")

doc.add_heading("5.6 Order Processing", level=2)
add_p("A user opens the order dialog from a watchlist action. The dialog receives the selected symbol, side, and market price. It accepts a positive integer quantity and a positive limit price, displays the estimated order value, and posts the validated values to the orders API. Mongoose applies schema constraints and records timestamps. The order-history route retrieves records in descending creation order.")

doc.add_heading("5.7 Fund Processing", level=2)
add_p("Fund balances use atomic MongoDB increments. For withdrawal, the update query includes a balance-greater-than-or-equal condition, so concurrent requests cannot reduce the balance below zero through a normal update. For deposit, the payment verification function reads the server-stored amount rather than trusting an amount returned by the browser. The PaymentOrder status changes from CREATED to PAID, and a previously paid order returns the current account without adding the amount again.")
add_table(["State", "Event", "Next state", "Balance action"], [
    ["CREATED", "Valid Razorpay signature", "PAID", "Credit stored order amount"],
    ["PAID", "Repeated verification", "PAID", "No additional credit"],
    ["CREATED", "Invalid signature", "CREATED", "Reject request"],
    ["Account active", "Valid withdrawal with sufficient balance", "Account active", "Atomic debit"],
], [1.3, 2.0, 1.3, 1.8])
add_caption("Figure 5.2 Fund transaction state transition")

doc.add_heading("5.8 Deployment", level=2)
add_p("The landing site and dashboard are linked to separate Vercel projects. Production builds execute react-scripts build, package static assets, and deploy serverless functions from the api directories. Rewrite rules direct client-side routes to index.html. Sensitive values are configured as deployment environment variables. The deployed landing site redirects authenticated users to the deployed dashboard address.")

chapter(6, "Testing and Results")
doc.add_heading("6.1 Testing Strategy", level=2)
add_p("Testing combined static inspection, syntax checks, production builds, API status checks, and manual workflow review. The backend entry file was checked with the Node.js syntax checker. Both React applications completed optimized production builds. The deployed landing page, dashboard, funds API, and authentication endpoint were checked for reachability. Input paths were reviewed for both successful and rejected values.")

doc.add_heading("6.2 Functional Test Cases", level=2)
add_table(["ID", "Test condition", "Expected result", "Observed result"], [
    ["TC-01", "Open public routes", "Requested page renders", "Pass"],
    ["TC-02", "Register valid new email", "Hashed user record and authenticated response", "Implemented"],
    ["TC-03", "Register duplicate email", "Conflict response directing user to login", "Implemented"],
    ["TC-04", "Login with invalid password", "Unauthorized response", "Implemented"],
    ["TC-05", "Submit fractional order quantity", "Validation error", "Pass by inspection"],
    ["TC-06", "Submit valid simulated order", "Order stored and returned", "Implemented"],
    ["TC-07", "Withdraw above available cash", "Request rejected", "Implemented"],
    ["TC-08", "Complete Razorpay test payment", "Signature verified and balance credited", "Implemented"],
    ["TC-09", "Repeat verified payment callback", "No duplicate credit", "Protected by state and transaction"],
    ["TC-10", "Build landing and dashboard", "Optimized bundles produced", "Pass"],
], [0.65, 2.15, 2.35, 1.35])
add_caption("Table 6.1 Functional test cases")

doc.add_heading("6.3 Build and Verification Results", level=2)
add_table(["Component", "Check", "Result"], [
    ["Backend", "Node syntax and package check script", "Passed"],
    ["Landing application", "Optimized React production build", "Passed with legacy accessibility warnings"],
    ["Dashboard application", "Optimized React production build", "Passed"],
    ["Landing deployment", "Login page HTTP reachability", "HTTP 200"],
    ["Authentication API", "Unsupported GET handling", "HTTP 405 as designed"],
    ["Dashboard deployment", "Root page reachability", "HTTP 200"],
    ["Funds API", "Read endpoint reachability", "HTTP 200"],
], [1.8, 3.3, 1.3])
add_caption("Table 6.2 Build and verification results")

doc.add_heading("6.4 Result Discussion", level=2)
add_p("The project achieves the intended academic demonstration. Users can move from a public product site to account registration or login and then explore a deployed dashboard. The dashboard displays investment data, records simulated orders, and manages a persistent virtual fund balance. Razorpay Test Checkout demonstrates an external payment provider without handling live money, and server-side HMAC verification prevents the browser from directly asserting that a payment succeeded.")
add_p("The results also reveal that a successful build is not equivalent to production readiness. Automated test discovery currently finds no meaningful suites. The landing build reports accessibility warnings from empty links. Security review finds that the dashboard does not yet validate the session created by the landing domain and that financial data is not partitioned by user. These findings are treated as engineering results rather than hidden defects and lead directly to the recommendations in Chapter 7.")

chapter(7, "Security Analysis Limitations and Risk Management")
doc.add_heading("7.1 Implemented Controls", level=2)
add_bullets([
    "Passwords are stored as bcrypt hashes rather than readable values.",
    "Database and Razorpay secrets are read from server environment variables.",
    "Mutable order and fund inputs are validated on the API rather than only in the browser.",
    "Razorpay signatures are recomputed with HMAC SHA-256 and compared using a timing-safe operation.",
    "A stored payment-order state and MongoDB transaction prevent normal duplicate crediting.",
    "Withdrawal updates include a sufficient-balance condition and apply an atomic decrement.",
    "JSON request size is restricted in the standalone Express service.",
])

doc.add_heading("7.2 Current Limitations", level=2)
add_p("The most important limitation is incomplete authorization. The session cookie is issued by the landing application, while the dashboard is hosted at a different domain and its APIs do not validate that session. Direct dashboard access is therefore possible. In addition, funds use a single demonstration account key, and orders, holdings, and positions do not carry a user owner. Multiple users would consequently share records. This behaviour is acceptable only for an explicitly labelled single-account academic demonstration.")
add_p("The payment workflow depends on the browser sending the Checkout result to the verification endpoint. A webhook and reconciliation process are required so that an interrupted browser session cannot leave a successful provider payment without a local transaction. Before live use, the server should also retrieve and confirm final payment capture status. Authentication endpoints need rate limiting, logout, session validation, password reset, and optional email verification. Direct dependencies and the deprecated Create React App toolchain require upgrades, and automated tests must cover authorization and concurrent financial operations.")

doc.add_heading("7.3 Risk and Mitigation Summary", level=2)
add_table(["Risk", "Impact", "Recommended mitigation"], [
    ["Unauthenticated dashboard API", "Unauthorised reading or modification", "Shared authentication service and mandatory API middleware"],
    ["Shared financial records", "Cross-user data exposure", "Add userId to every entity and filter every query"],
    ["Client-only payment completion", "Paid transaction may not be credited", "Signed Razorpay webhooks and reconciliation job"],
    ["No login rate limit", "Credential guessing and resource abuse", "IP and account-based throttling with monitoring"],
    ["Outdated dependencies", "Known security and reliability defects", "Controlled package upgrades and regression testing"],
    ["No automated tests", "Undetected regression", "Unit, API integration and end-to-end suites in CI"],
], [1.65, 1.75, 3.0])
add_caption("Table 7.1 Risk and mitigation summary")

doc.add_heading("7.4 Ethical and Regulatory Position", level=2)
add_p("Stashflow must remain clearly identified as a learning prototype until the controls described above are implemented and independently reviewed. It should not claim to execute real securities trades, hold customer funds, provide investment advice, or satisfy KYC and regulatory obligations. Test data should avoid real personal financial information. Any future live version would require legal, security, privacy, payment, and financial-services compliance review beyond the scope of this academic project.")

chapter(8, "Conclusion and Future Scope")
doc.add_heading("8.1 Conclusion", level=2)
add_p("Stashflow demonstrates the integration of a responsive React interface, account registration and login, persistent MongoDB data, serverless API functions, chart-based portfolio presentation, simulated order processing, virtual fund management, and Razorpay Test Checkout. The project transforms a static trading clone into a deployed full-stack prototype with meaningful validation and payment-integrity controls.")
add_p("The implementation provides a useful academic case study because it includes both successful engineering decisions and clearly observable limitations. Password hashing, server-side secret handling, atomic balance updates, signature verification, and idempotent payment crediting provide a sound foundation. At the same time, the review shows that authentication must be enforced across the dashboard, records must be owned by individual users, webhooks must support payment reconciliation, and testing and dependencies must be strengthened before any real financial use.")

doc.add_heading("8.2 Future Scope", level=2)
add_bullets([
    "Place the landing site and dashboard behind one authentication boundary, or implement a secure shared identity provider.",
    "Introduce per-user portfolio, order, payment, and fund-account ownership with migration of demonstration records.",
    "Add logout, session renewal and revocation, password reset, email verification, multi-factor authentication, and rate limiting.",
    "Process Razorpay webhooks, verify captured status, reconcile missed callbacks, and build an administrative payment ledger.",
    "Connect a licensed market-data source for delayed or real-time quotes and record quote timestamps.",
    "Add portfolio analytics, realised and unrealised profit, allocation targets, downloadable statements, and tax summaries.",
    "Replace Create React App with a maintained build tool and update vulnerable dependencies through tested releases.",
    "Create unit, integration, accessibility, security, and end-to-end test suites executed on every change.",
    "Add audit logs, monitoring, alerts, backups, disaster recovery, and documented data-retention controls.",
    "Conduct usability testing and accessibility conformance review for keyboard, screen-reader, contrast, and mobile use.",
])

chapter(9, "References")
refs = [
    "[1] Dr A P J Abdul Kalam Technical University, Uttar Pradesh, official university portal. Available: https://aktu.ac.in/",
    "[2] React, “React Documentation.” Available: https://react.dev/",
    "[3] React Router, “React Router Documentation.” Available: https://reactrouter.com/",
    "[4] MongoDB, “MongoDB Documentation.” Available: https://www.mongodb.com/docs/",
    "[5] Mongoose, “Mongoose Documentation.” Available: https://mongoosejs.com/docs/",
    "[6] Node.js, “Node.js Documentation.” Available: https://nodejs.org/docs/latest/api/",
    "[7] Express.js, “Express Web Framework Documentation.” Available: https://expressjs.com/",
    "[8] Razorpay, “Standard Checkout Integration Steps.” Available: https://razorpay.com/docs/payments/payment-gateway/web-integration/standard/integration-steps/",
    "[9] Razorpay, “API Keys on the Razorpay Dashboard.” Available: https://razorpay.com/docs/payments/dashboard/account-settings/api-keys/",
    "[10] Vercel, “Vercel Documentation.” Available: https://vercel.com/docs",
    "[11] OWASP Foundation, “Authentication Cheat Sheet.” Available: https://cheatsheetseries.owasp.org/cheatsheets/Authentication_Cheat_Sheet.html",
    "[12] OWASP Foundation, “Session Management Cheat Sheet.” Available: https://cheatsheetseries.owasp.org/cheatsheets/Session_Management_Cheat_Sheet.html",
    "[13] OWASP Foundation, “Transaction Authorization Cheat Sheet.” Available: https://cheatsheetseries.owasp.org/cheatsheets/Transaction_Authorization_Cheat_Sheet.html",
    "[14] Chart.js, “Chart.js Documentation.” Available: https://www.chartjs.org/docs/latest/",
    "[15] npm, “npm Documentation.” Available: https://docs.npmjs.com/",
]
for ref in refs:
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(0.25); p.paragraph_format.first_line_indent = Inches(-0.25); p.paragraph_format.line_spacing = 1.25
    font(p.add_run(ref + " Accessed October 2026."), 11)

chapter("A", "Installation and Execution")
doc.add_heading("A.1 Prerequisites", level=2)
add_bullets(["Node.js and npm", "MongoDB connection string", "Razorpay Test Mode key identifier and secret", "Modern browser", "Vercel account for cloud deployment"])
doc.add_heading("A.2 Local Configuration", level=2)
add_p("Create local environment files that are excluded from version control. Configure MONGO_URL for database access. Configure SESSION_SECRET for the authentication API. Configure RAZORPAY_KEY_ID and RAZORPAY_KEY_SECRET only for the dashboard server functions. Public browser code must never contain the Razorpay secret or MongoDB connection string.")
doc.add_heading("A.3 Commands", level=2)
add_table(["Working directory", "Command", "Purpose"], [
    ["frontend", "npm install", "Install landing and authentication dependencies"],
    ["frontend", "npm start", "Run public application in development"],
    ["frontend", "npm run build", "Create optimized landing build"],
    ["dashboard", "npm install", "Install dashboard and payment dependencies"],
    ["dashboard", "npm start", "Run dashboard in development"],
    ["dashboard", "npm run build", "Create optimized dashboard build"],
    ["backend", "npm install", "Install standalone Express dependencies"],
    ["backend", "npm run dev", "Run local Express API with file watching"],
], [1.4, 1.7, 3.3])

chapter("B", "API Summary")
add_table(["Method", "Endpoint", "Responsibility"], [
    ["POST", "/api/auth", "Register or authenticate a user"],
    ["GET", "/api/holdings", "Return holdings records"],
    ["GET", "/api/positions", "Return positions records"],
    ["GET", "/api/orders", "Return stored simulated orders"],
    ["POST", "/api/orders", "Validate and store a simulated order"],
    ["GET", "/api/funds", "Return balance and recent transactions"],
    ["POST", "/api/funds", "Process a virtual bank withdrawal"],
    ["POST", "/api/razorpay-order", "Create a Razorpay test order"],
    ["POST", "/api/razorpay-verify", "Verify the payment signature and credit funds"],
], [0.9, 2.3, 3.2])

chapter("C", "Sample Viva Questions")
questions = [
    ("Why is the Razorpay secret not sent to React?", "It is a server credential. Exposing it would allow unauthorised API calls and signature generation."),
    ("Why is HMAC used during payment verification?", "It proves that the response fields correspond to a party holding the shared secret and detects modification."),
    ("How is duplicate payment credit prevented?", "A unique payment order stores CREATED or PAID state, and the status change and balance update execute in one database transaction."),
    ("Why use bcrypt for passwords?", "Bcrypt is deliberately slow, salted, and designed for password hashing, which increases the cost of offline guessing."),
    ("Why use atomic update operators for balance changes?", "They avoid a separate read-modify-write window and reduce lost updates during concurrent requests."),
    ("What is the main current limitation?", "The dashboard APIs do not yet enforce the session or isolate records per user, so the deployment remains an academic prototype."),
    ("What is the role of React Router?", "It maps browser paths to components without a complete page reload and supports separate dashboard views."),
    ("Why are webhooks needed?", "They allow the server to learn about payment events even when the user closes the browser or a client callback fails."),
]
for i, (q, a) in enumerate(questions, 1):
    add_p(f"Q{i}. {q}", bold_lead=f"Q{i}. ")
    add_p(f"Answer: {a}", bold_lead="Answer: ")

# Footer label and metadata
for section in doc.sections:
    section.header_distance = Inches(0.4); section.footer_distance = Inches(0.45)
    section.header.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    for child in list(hp._p):
        hp._p.remove(child)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(hp.add_run("Stashflow Project Report"), 9, False, color="555555")

settings = doc.settings.element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields"); settings.append(update)
update.set(qn("w:val"), "true")

doc.core_properties.title = "Stashflow Web Based Investment and Trading Dashboard"
doc.core_properties.subject = "AKTU B Tech Project Report"
doc.core_properties.author = "Project Team"
doc.core_properties.keywords = "AKTU, B Tech, Stashflow, React, MongoDB, Razorpay"
doc.save(OUT)
print(OUT)
