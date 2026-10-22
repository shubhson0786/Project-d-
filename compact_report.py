from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

root=Path(__file__).parent
doc=Document(root/'Stashflow_AKTU_Project_Report_Completed.docx')
normal=doc.styles['Normal']
normal.font.name='Times New Roman'
normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.5
for name, size in [('Heading 1',14),('Heading 2',12),('Heading 3',12)]:
    style=doc.styles[name]
    style.font.name='Times New Roman'
    style.font.size=Pt(size)
    style.font.bold=True
    style.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
for section in doc.sections:
    section.page_width=Inches(8.27)
    section.page_height=Inches(11.69)
    section.top_margin=Inches(1)
    section.bottom_margin=Inches(1)
    section.left_margin=Inches(1.25)
    section.right_margin=Inches(1.25)
    for paragraph in section.footer.paragraphs:
        paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.font.name='Times New Roman'
            run.font.size=Pt(12)
# Let chapters flow naturally rather than leaving nearly empty continuation pages.
for p in doc.paragraphs:
    if p.text.startswith('Note: In Microsoft Word, select the table and choose Update Field'):
        p._element.getparent().remove(p._element)
        continue
    if p.style.name == 'Heading 1':
        p.paragraph_format.page_break_before=False
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_before=Pt(18)
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name='Times New Roman'
            run.font.size=Pt(14)
            run.bold=True
    if p.style.name in ('Heading 2','Heading 3'):
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name='Times New Roman'
            run.font.size=Pt(12)
            run.bold=True
    if not p.text.strip() and not p._p.xpath('.//w:br | .//w:sectPr | .//w:fldChar'):
        p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.space_after=Pt(0)
        p.paragraph_format.line_spacing=Pt(1)
        for r in p.runs:r.font.size=Pt(1)
for section in list(doc.sections)[1:]:
    section.different_first_page_header_footer=False
    section.start_type=2
doc.save(root/'Stashflow_AKTU_Report_Corrected.docx')
print(root/'Stashflow_AKTU_Report_Corrected.docx')
